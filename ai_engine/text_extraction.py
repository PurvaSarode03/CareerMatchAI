"""Extract clean text from PDF / DOCX / TXT files."""
import io
import re

from docx import Document
from pypdf import PdfReader


class ExtractionError(Exception):
    pass


_BULLETS = "•●▪■◦◆○➢➤►–·\uf0b7\uf0a7"


def _normalize(text: str) -> str:
    text = text.replace("\r", "\n").replace("\x00", "")
    text = re.sub(f"[{re.escape(_BULLETS)}]", "• ", text)
    text = re.sub(r"[ \t\u00a0]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.split("\n")).strip()


def extract_pdf(fileobj) -> str:
    try:
        reader = PdfReader(fileobj)
        if reader.is_encrypted:
            raise ExtractionError("This PDF is password protected.")
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError("Could not read this PDF. It may be corrupted.") from exc
    return text


def extract_docx(fileobj) -> str:
    try:
        doc = Document(fileobj)
    except Exception as exc:
        raise ExtractionError("Could not read this DOCX file.") from exc
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # resumes often use tables for layout
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if cell._tc not in seen:
                    seen.add(cell._tc)
                    parts.append(cell.text)
    return "\n".join(parts)


def extract_text(fileobj, filename: str) -> str:
    """Return normalized text. Raises ExtractionError for unreadable/empty files."""
    name = filename.lower()
    if hasattr(fileobj, "seek"):
        fileobj.seek(0)
    data = io.BytesIO(fileobj.read())
    if name.endswith(".pdf"):
        raw = extract_pdf(data)
    elif name.endswith(".docx"):
        raw = extract_docx(data)
    elif name.endswith(".txt"):
        raw = data.getvalue().decode("utf-8", errors="ignore")
    else:
        raise ExtractionError("Unsupported file type.")
    text = _normalize(raw)
    if len(text) < 40:
        raise ExtractionError("No readable text found. Scanned/image-only files are not supported — upload a text-based PDF or DOCX.")
    return text
